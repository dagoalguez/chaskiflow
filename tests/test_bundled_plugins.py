"""Pruebas de los plugins incluidos (http_request, export_csv, export_xlsx) y del CLI."""

import csv
import json
import shutil
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import run_workflow  # noqa: E402
from chaskiflow.engine import Engine  # noqa: E402
from chaskiflow.plugin_loader import PluginRegistry  # noqa: E402
from tests.helpers import PLUGINS, ROOT, LocalServer  # noqa: E402

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.engine = Engine(PluginRegistry([PLUGINS]), workdir_root=cls.tmp / "runs")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def run_node(self, typ, config):
        w = {"name": "t", "nodes": [{"id": "A", "label": "A", "type": typ, "config": config}],
             "edges": []}
        return self.engine.run(w).nodes["A"]


class HttpRequestTests(Base):
    def test_get_json(self):
        with LocalServer() as s:
            r = self.run_node("http_request", {"url": s.base + "/json"})
        self.assertEqual(r["status"], "ok", r["error"])
        res = r["result"]
        self.assertEqual((res["status"], res["ok"], res["json"]), (200, True, {"a": 1, "texto": "ñandú"}))
        self.assertIn("content-type", res["headers"])

    def test_get_text(self):
        with LocalServer() as s:
            res = self.run_node("http_request", {"url": s.base + "/text"})["result"]
        self.assertEqual(res["text"], "hola texto")
        self.assertIsNone(res["json"])

    def test_http_error_fails_by_default(self):
        with LocalServer() as s:
            r = self.run_node("http_request", {"url": s.base + "/status/404"})
        self.assertEqual(r["status"], "error")
        self.assertIn("HTTP 404", r["error"])

    def test_http_error_can_be_tolerated(self):
        with LocalServer() as s:
            r = self.run_node("http_request", {"url": s.base + "/status/404",
                                               "fail_on_http_error": False})
        self.assertEqual(r["status"], "ok")
        self.assertEqual((r["result"]["status"], r["result"]["ok"], r["result"]["text"]),
                         (404, False, "no existe"))

    def test_post_dict_body_is_sent_as_json(self):
        with LocalServer() as s:
            res = self.run_node("http_request", {
                "method": "POST", "url": s.base + "/echo",
                "body": {"mensaje": "hola ñ", "n": [1, 2]},
                "headers": {"Authorization": "Bearer xyz"}})["result"]
        self.assertEqual(res["json"]["received"], {"mensaje": "hola ñ", "n": [1, 2]})
        self.assertEqual(res["json"]["auth"], "Bearer xyz")
        self.assertIn("application/json", res["json"]["ctype"])

    def test_headers_as_json_text(self):
        with LocalServer() as s:
            res = self.run_node("http_request", {
                "method": "POST", "url": s.base + "/echo", "body": "texto plano",
                "headers": '{"Authorization": "Bearer abc"}'})["result"]
        self.assertEqual(res["json"]["received"], "texto plano")
        self.assertEqual(res["json"]["auth"], "Bearer abc")

    def test_timeout(self):
        with LocalServer() as s:
            r = self.run_node("http_request", {"url": s.base + "/slow", "timeout": 1})
        self.assertEqual(r["status"], "error")
        self.assertIn("Tiempo de espera agotado", r["error"])

    def test_connection_refused(self):
        r = self.run_node("http_request", {"url": "http://127.0.0.1:1/x", "timeout": 3})
        self.assertEqual(r["status"], "error")
        self.assertIn("No se pudo conectar", r["error"])

    def test_rejects_non_http_urls(self):
        r = self.run_node("http_request", {"url": "file:///etc/passwd"})
        self.assertEqual(r["status"], "error")
        self.assertIn("http", r["error"])

    def test_user_agent(self):
        with LocalServer() as s:
            default = self.run_node("http_request", {"url": s.base + "/ua"})["result"]["json"]["ua"]
            custom = self.run_node("http_request", {"url": s.base + "/ua",
                                                    "user_agent": "MiBot/1.0"})["result"]["json"]["ua"]
        self.assertIn("ChaskiFlow", default)
        self.assertEqual(custom, "MiBot/1.0")


ROWS = [
    {"medio": "RPP", "titulo": "Evasión; con \"comillas\"\ny salto", "n": 3, "ok": True, "extra": None},
    {"medio": "Gestión", "titulo": "=1+1", "n": 2.5, "ok": False, "datos": {"a": [1]}},
    {"medio": "Infobae", "titulo": "-cmd", "n": 0, "ok": True},
    {"medio": "Perú21", "titulo": "ñandú café", "n": -4, "ok": True},
]


class ExportCsvTests(Base):
    def read(self, path):
        raw = Path(path).read_bytes()
        text = raw.decode("utf-8-sig")
        return raw, list(csv.reader(text.splitlines(True) and __import__("io").StringIO(text),
                                    delimiter=";"))

    def test_basic_export(self):
        r = self.run_node("export_csv", {"data": ROWS, "output_dir": str(self.tmp / "csv1"),
                                         "filename": "rep_%Y"})
        self.assertEqual(r["status"], "ok", r["error"])
        res = r["result"]
        raw, table = self.read(res["file_path"])
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))  # BOM para Excel
        self.assertEqual(table[0], ["medio", "titulo", "n", "ok", "extra", "datos"])
        self.assertEqual(table[1][1], 'Evasión; con "comillas"\ny salto')
        self.assertEqual(table[1][2:5], ["3", "true", ""])
        self.assertEqual(table[2][5], '{"a": [1]}')
        self.assertEqual(res["rows"], 4)
        self.assertTrue(res["filename"].endswith(".csv"))
        self.assertEqual(res["file_paths"], [res["file_path"]])

    def test_formula_injection_neutralized(self):
        r = self.run_node("export_csv", {"data": ROWS, "output_dir": str(self.tmp / "csv2")})
        _, table = self.read(r["result"]["file_path"])
        self.assertEqual(table[2][1], "'=1+1")
        self.assertEqual(table[3][1], "'-cmd")
        self.assertEqual(table[4][2], "-4")  # los números no se tocan
        r = self.run_node("export_csv", {"data": ROWS, "output_dir": str(self.tmp / "csv3"),
                                         "escape_formulas": False})
        _, table = self.read(r["result"]["file_path"])
        self.assertEqual(table[2][1], "=1+1")

    def test_columns_rename_delimiter_encoding(self):
        r = self.run_node("export_csv", {
            "data": ROWS, "output_dir": str(self.tmp / "csv4"), "columns": "titulo, medio",
            "rename": {"titulo": "Título", "medio": "Medio"}, "delimiter": "|",
            "encoding": "utf-8"})
        raw = Path(r["result"]["file_path"]).read_bytes()
        self.assertFalse(raw.startswith(b"\xef\xbb\xbf"))
        self.assertTrue(raw.decode("utf-8").startswith("Título|Medio\n") or
                        raw.decode("utf-8").startswith("Título|Medio\r\n"))

    def test_overwrite_false_numbers_files(self):
        cfg = {"data": [{"a": 1}], "output_dir": str(self.tmp / "csv5"), "filename": "x",
               "overwrite": False}
        names = [self.run_node("export_csv", cfg)["result"]["filename"] for _ in range(3)]
        self.assertEqual(names, ["x.csv", "x_1.csv", "x_2.csv"])
        cfg["overwrite"] = True
        self.assertEqual(self.run_node("export_csv", cfg)["result"]["filename"], "x.csv")

    def test_filename_sanitized_and_strftime(self):
        r = self.run_node("export_csv", {"data": [{"a": 1}], "output_dir": str(self.tmp / "csv6"),
                                         "filename": "a:b*c_%Y"})
        self.assertRegex(r["result"]["filename"], r"^a_b_c_\d{4}\.csv$")

    def test_max_cell_chars(self):
        r = self.run_node("export_csv", {"data": [{"t": "x" * 100}], "max_cell_chars": 10,
                                         "output_dir": str(self.tmp / "csv7")})
        _, table = self.read(r["result"]["file_path"])
        self.assertEqual(table[1][0], "x" * 10)

    def test_empty_data_still_writes_header_only(self):
        r = self.run_node("export_csv", {"data": [], "columns": "a,b",
                                         "output_dir": str(self.tmp / "csv8")})
        self.assertEqual(r["status"], "ok")
        _, table = self.read(r["result"]["file_path"])
        self.assertEqual(table, [["a", "b"]])

    def test_bad_data_gives_clear_error(self):
        r = self.run_node("export_csv", {"data": "no es lista", "output_dir": str(self.tmp / "c9")})
        self.assertEqual(r["status"], "error")
        self.assertIn("lista", r["error"])
        r = self.run_node("export_csv", {"data": [1, 2], "output_dir": str(self.tmp / "c9")})
        self.assertIn("objetos", r["error"])


class ExportXlsxTests(Base):
    def sheet_rows(self, path):
        with zipfile.ZipFile(path) as z:
            sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            wb = ET.fromstring(z.read("xl/workbook.xml"))
        rows = []
        for row in sheet.find("m:sheetData", NS):
            cells = {}
            for c in row:
                col = "".join(ch for ch in c.get("r") if ch.isalpha())
                if c.get("t") == "inlineStr":
                    val = "".join(t.text or "" for t in c.iter("{%s}t" % NS["m"]))
                else:
                    v = c.find("m:v", NS)
                    val = v.text if v is not None else None
                    if c.get("t") == "b":
                        val = bool(int(val))
                    elif val is not None:
                        val = float(val) if "." in val else int(val)
                cells[col] = val
            rows.append(cells)
        return rows, wb.find("m:sheets", NS)[0].get("name"), sheet

    def test_structure_and_values(self):
        r = self.run_node("export_xlsx", {"data": ROWS, "output_dir": str(self.tmp / "x1"),
                                          "sheet_name": "Noticias"})
        self.assertEqual(r["status"], "ok", r["error"])
        rows, title, sheet = self.sheet_rows(r["result"]["file_path"])
        self.assertEqual(title, "Noticias")
        self.assertEqual(rows[0], {"A": "medio", "B": "titulo", "C": "n", "D": "ok", "E": "extra",
                                   "F": "datos"})
        self.assertEqual(rows[1]["B"], 'Evasión; con "comillas"\ny salto')
        self.assertEqual((rows[1]["C"], rows[1]["D"]), (3, True))
        self.assertNotIn("E", rows[1])  # None -> celda vacía
        self.assertEqual(rows[2]["C"], 2.5)
        self.assertEqual(rows[4]["C"], -4)
        self.assertEqual(rows[2]["F"], '{"a": [1]}')
        self.assertEqual(rows[2]["B"], "=1+1")  # guardado como TEXTO, no como fórmula
        self.assertIsNone(sheet.find(".//m:f", NS))
        self.assertIsNotNone(sheet.find("m:autoFilter", NS))
        self.assertIsNotNone(sheet.find(".//m:pane", NS))

    def test_header_is_bold_style(self):
        r = self.run_node("export_xlsx", {"data": [{"a": 1}], "output_dir": str(self.tmp / "x2")})
        _, _, sheet = self.sheet_rows(r["result"]["file_path"])
        first = sheet.find("m:sheetData", NS)[0][0]
        self.assertEqual(first.get("s"), "1")

    def test_illegal_chars_and_limits(self):
        r = self.run_node("export_xlsx", {
            "data": [{"t": "a\x00b\x07c\u0001d" + "z" * 40000}], "output_dir": str(self.tmp / "x3")})
        self.assertEqual(r["status"], "ok", r["error"])
        rows, _, _ = self.sheet_rows(r["result"]["file_path"])
        self.assertTrue(rows[1]["A"].startswith("abcd"))
        self.assertEqual(len(rows[1]["A"]), 32000)  # max_cell_chars por defecto

    def test_sheet_name_sanitized(self):
        r = self.run_node("export_xlsx", {"data": [{"a": 1}], "output_dir": str(self.tmp / "x4"),
                                          "sheet_name": "a/b:c?" + "x" * 40})
        _, title, _ = self.sheet_rows(r["result"]["file_path"])
        self.assertEqual(title, ("a_b_c_" + "x" * 40)[:31])

    def test_opens_with_openpyxl_if_available(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no está instalado (solo se usa para verificar en desarrollo)")
        import warnings
        r = self.run_node("export_xlsx", {"data": ROWS, "output_dir": str(self.tmp / "x5"),
                                          "sheet_name": "Datos"})
        with warnings.catch_warnings():
            warnings.simplefilter("error")  # cualquier aviso de formato debe fallar
            wb = openpyxl.load_workbook(r["result"]["file_path"])
        ws = wb["Datos"]
        self.assertEqual((ws.max_row, ws.max_column), (5, 6))
        self.assertEqual(ws["B3"].value, "=1+1")
        self.assertEqual(ws["B3"].data_type, "s")  # texto, no fórmula
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["C2"].value, 3)
        self.assertIs(ws["D2"].value, True)

    def test_many_rows(self):
        data = [{"i": i, "t": "fila %d" % i} for i in range(20000)]
        r = self.run_node("export_xlsx", {"data": data, "output_dir": str(self.tmp / "x6")})
        self.assertEqual(r["status"], "ok", r["error"])
        self.assertEqual(r["result"]["rows"], 20000)


class CliTests(Base):
    def test_example_workflow_end_to_end(self):
        out = self.tmp / "cli_out"
        code = run_workflow.main([str(ROOT / "examples" / "hola_reporte.json"),
                                  "--var", "carpeta=%s" % out, "--workdir", str(self.tmp / "w")])
        self.assertEqual(code, 0)
        self.assertEqual(len(list(out.glob("*.csv"))), 1)
        self.assertEqual(len(list(out.glob("*.xlsx"))), 1)

    def test_validate_reports_errors(self):
        bad = self.tmp / "bad.json"
        bad.write_text(json.dumps({"nodes": [{"id": "A", "label": "A", "type": "nada"}]}))
        self.assertEqual(run_workflow.main([str(bad), "--validate"]), 2)

    def test_list_plugins(self):
        self.assertEqual(run_workflow.main(["--list-plugins"]), 0)


if __name__ == "__main__":
    unittest.main()
